import re
from collections import Counter
from docx import Document


def normalize_text(text):
    """Normalize whitespace for reliable text comparisons."""
    return re.sub(r"\s+", " ", text).strip()


def extract_document_structure(docx_path):
    """Extract paragraphs, headings, tables and duplicate paragraph counts."""
    doc = Document(str(docx_path))

    paragraphs = [
        {
            "text": normalize_text(p.text),
            "style": p.style.name
        }
        for p in doc.paragraphs
        if normalize_text(p.text)
    ]

    headings = [
        item for item in paragraphs
        if item["style"].lower().startswith("heading")
    ]

    paragraph_counts = Counter(
        item["text"].casefold() for item in paragraphs
    )

    duplicates = [
        text for text, count in paragraph_counts.items()
        if count > 1
    ]

    tables = [
        [
            [cell.text.strip() for cell in row.cells]
            for row in table.rows
        ]
        for table in doc.tables
    ]

    return {
        "paragraph_count": len(paragraphs),
        "heading_count": len(headings),
        "table_count": len(tables),
        "paragraphs": paragraphs,
        "headings": headings,
        "tables": tables,
        "duplicate_paragraph_count": len(duplicates),
        "duplicate_paragraphs": duplicates
    }


def check_main_heading_numbering(headings):
    """Check whether Heading 1 numbering is sequential from 1."""
    main_headings = [
        item["text"] for item in headings
        if item["style"] == "Heading 1"
    ]

    numbers = []
    for heading in main_headings:
        match = re.match(r"^\s*(\d+)\.", heading)
        if not match:
            return {
                "status": "REVIEW",
                "numbers": numbers,
                "details": f"Unnumbered main heading: {heading}"
            }
        numbers.append(int(match.group(1)))

    expected = list(range(1, len(main_headings) + 1))
    passed = numbers == expected

    return {
        "status": "PASS" if passed else "REVIEW",
        "numbers": numbers,
        "details": (
            "Sequential numbering verified."
            if passed else
            "Main heading numbering requires review."
        )
    }
